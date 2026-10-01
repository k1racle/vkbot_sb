"""Read-only XLSX exports from the current project's saved client directory."""

import re
from datetime import datetime, timezone
from tempfile import SpooledTemporaryFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import clients, db, projects

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XML_UNSAFE = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
HEADERS = (
    "ID VK",
    "Имя",
    "Фамилия",
    "Ссылка VK",
    "Телефон",
    "Источник телефона",
    "Ссылка на фото",
    "Последняя отправка бота (UTC)",
    "Последнее входящее (UTC)",
    "Сообщения VK: последнее известное разрешение",
    "Отписался от рассылок",
    "Профиль недоступен",
    "Обновление профиля (UTC)",
    "Добавлен (UTC)",
    "Email",
    "Мессенджер",
    "Контакты: дата и источник",
)


def cells(sheet, values, *, heading=False):
    result = []
    for value in values:
        if isinstance(value, str):
            value = _XML_UNSAFE.sub("", value)[:32767]
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            # Never interpret user names/phones as formulae or Excel errors.
            cell.data_type = "s"
        elif isinstance(value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
        if heading:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="227E68")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        result.append(cell)
    return result


def client_workbook(search="", contacted=False):
    """Spool the finished ZIP, stream DB rows, never call VK or fetch photos."""
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    workbook = Workbook(write_only=True)
    try:
        project = projects.current_project.get()
        info = workbook.create_sheet("О выгрузке")
        info.column_dimensions["A"].width = 28
        info.column_dimensions["B"].width = 80
        for values in (
            ("Проект", project.name if project else "Исходный проект"),
            ("ID проекта", str(project.id) if project else ""),
            ("ID сообщества", str(project.group_id) if project else ""),
            ("Выгружено (UTC)", datetime.now(timezone.utc).replace(tzinfo=None)),
            ("Поиск", search),
            ("Кому бот писал", "Да" if contacted else "Все клиенты"),
            (
                "Состав",
                "Все найденные клиенты, не только текущая страница. Данные из базы, без запроса к VK.",
            ),
            (
                "Фото",
                "Ссылки, без скачивания изображений. Отсутствующие данные оставлены пустыми.",
            ),
        ):
            info.append(cells(info, values))

        def new_sheet(number):
            sheet = workbook.create_sheet(
                "Клиенты" if number == 1 else f"Клиенты {number}"
            )
            sheet.freeze_panes = "A2"
            sheet.row_dimensions[1].height = 45
            for index in range(1, len(HEADERS) + 1):
                sheet.column_dimensions[get_column_letter(index)].width = (
                    26 if index != 7 else 48
                )
            sheet.append(cells(sheet, HEADERS, heading=True))
            return sheet

        sheet = new_sheet(1)
        workbook.move_sheet(info.title, offset=1)
        count, number = 1, 1
        with db.SessionLocal() as session:
            query = clients.list_query(session, search, contacted).order_by(
                db.Client.user_id.desc()
            )
            for client in query.yield_per(500):
                if count == 1048576:
                    sheet.auto_filter.ref = (
                        f"A1:{get_column_letter(len(HEADERS))}{count}"
                    )
                    number += 1
                    sheet, count = new_sheet(number), 1
                sheet.append(
                    cells(
                        sheet,
                        (
                            str(client.user_id),
                            client.first_name,
                            client.last_name,
                            f"https://vk.ru/id{client.user_id}",
                            client.phone,
                            {"vk": "VK", "dialog": "Диалог"}.get(
                                client.phone_source, ""
                            ),
                            clients.safe_photo(client.photo_url),
                            client.bot_contacted_at,
                            client.last_incoming_at,
                            "Неизвестно"
                            if client.messages_allowed is None
                            else "Да"
                            if client.messages_allowed
                            else "Нет",
                            "Да" if client.unsubscribed else "Нет",
                            "Да" if client.deactivated else "Нет",
                            client.profile_updated_at,
                            client.created_at,
                            client.email,
                            client.messenger,
                            "; ".join(
                                f"{kind}: {details.get('date', '')} {details.get('scenario', '')}"
                                for kind, details in (
                                    client.contact_details or {}
                                ).items()
                            ),
                        ),
                    )
                )
                count += 1
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{count}"
        workbook.save(output)
        output.seek(0)
        return output
    except BaseException:
        output.close()
        raise
    finally:
        workbook.close()


def chunks(file):
    try:
        while data := file.read(65536):
            yield data
    finally:
        file.close()
