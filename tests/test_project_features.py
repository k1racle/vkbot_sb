"""Campaign keywords, isolated XLSX and video invitations: all VK is mocked."""

import asyncio
from datetime import datetime
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from test_projects import (
    project_env as project_env,
    two_projects as two_projects,
    callback,
    comment,
    configure,
    db,
    login,
    message,
    prefix,
    projects,
    session_for,
)
from app import config, vk_api
from app.comments import matches_plus_words, normalize_comment


def campaign(project, **fields):
    with session_for(project) as session:
        row = db.Campaign(
            post_id=0,
            title="Подарок",
            promo_code="GIFT",
            promo_message="Код {promo_code}",
            **fields,
        )
        session.add(row)
        session.commit()
        return row.id


def video(project, number=1, content="хочу подарок"):
    payload = comment(project, number)
    payload["type"] = "video_comment_new"
    payload["object"].pop("post_id")
    payload["object"].update(
        video_id=55, video_owner_id=-project.group_id, text=content
    )
    return payload


def status(project):
    with session_for(project) as session:
        return (
            session.query(db.ProcessedComment)
            .order_by(db.ProcessedComment.id.desc())
            .first()
            .status
        )


@pytest.mark.parametrize(
    "words,body,matched",
    [
        ("", "Любое сообщение", True),
        (" , \n ", "", True),
        ("подарок,хочу", "ХОЧУ!", True),
        ("подарок\nхочу", "(подарок)", True),
        ("крем", "крема", False),
        ("кот", "бойкот", False),
        ("кот", "котик", False),
        ("крем", "крем", True),
        ("хочу подарок", "Хочу\n   подарок!", True),
        ("всё", "Все", True),
        (".*", "другое", False),
        (".*", " .* ", True),
        ("🎁", "Подарок 🎁!", True),
    ],
)
def test_whole_plus_words_and_phrases(words, body, matched):
    assert bool(matches_plus_words(body, words)) is matched


@pytest.mark.parametrize("kind", ["wall", "video"])
@pytest.mark.parametrize("mode", ["direct", "chat_invite"])
def test_plus_filter_before_any_vk_and_stop_wins(two_projects, kind, mode):
    env, first, _ = two_projects
    campaign(
        first,
        plus_words="подарок",
        stop_words="спам",
        delivery_mode=mode,
        one_promo_per_user=False,
    )
    build = video if kind == "video" else comment
    for number, body, expected in [
        (1, "Привет", "plus_word_missing"),
        (2, "подарок спам", "stop_word"),
    ]:
        payload = build(first, number)
        payload["object"]["text"] = body
        assert callback(env.client, first, payload).text == "ok"
        assert status(first) == expected
    assert env.network.calls == []
    payload = build(first, 3)
    payload["object"]["text"] = "ПОДАРОК!"
    callback(env.client, first, payload)
    assert status(first) in {"sent", "waiting_chat", "video_invited_dm"}


def test_plus_admin_edit_clear_legacy_form_and_projects(two_projects):
    env, first, second = two_projects
    data = dict(
        title="Акция",
        promo_code="GIFT",
        promo_message="Код",
        shop_url="https://example.org",
        enabled="1",
        plus_words="хочу, подарок",
    )
    response = env.client.post(prefix(first) + "/campaigns", data=data)
    assert response.status_code == 200 and "campaign_saved=1" in str(response.url)
    with session_for(first) as session:
        row = session.query(db.Campaign).one()
        assert row.plus_words == "хочу, подарок"
        ident = row.id
    with session_for(second) as session:
        assert session.query(db.Campaign).count() == 0
    data["campaign_id"] = str(ident)
    data.pop("plus_words")
    env.client.post(prefix(first) + "/campaigns", data=data)
    with session_for(first) as session:
        assert session.get(db.Campaign, ident).plus_words == "хочу, подарок"
    data["plus_words"] = ""
    env.client.post(prefix(first) + "/campaigns", data=data)
    with session_for(first) as session:
        assert session.get(db.Campaign, ident).plus_words == ""
    assert not env.network.calls


def workbook(response):
    assert response.status_code == 200, (
        response.text[:300] if response.status_code != 200 else ""
    )
    assert "spreadsheetml" in response.headers["content-type"]
    assert "attachment;" in response.headers["content-disposition"]
    assert "no-store" in response.headers["cache-control"]
    return load_workbook(BytesIO(response.content))


def test_excel_all_rows_filters_types_and_group_isolation(two_projects):
    env, first, second = two_projects
    when = datetime(2026, 9, 28, 12, 34, 56)
    with session_for(first) as session:
        for ident in range(1, 66):
            session.add(
                db.Client(
                    user_id=ident,
                    first_name="Анна",
                    last_name="Первая",
                    bot_contacted_at=when if ident % 2 else None,
                )
            )
        session.add(
            db.Client(
                user_id=99,
                first_name='=HYPERLINK("https://evil.invalid")',
                last_name="@SUM(1)",
                phone="+79990000000",
                phone_source="dialog",
                photo_url="https://example.org/photo.jpg",
                messages_allowed=False,
                unsubscribed=True,
                last_incoming_at=when,
            )
        )
        session.add(db.Client(user_id=100, first_name="Иван\x00Пётр", last_name="#N/A"))
        session.commit()
    with session_for(second) as session:
        session.add(db.Client(user_id=99, first_name="Другая группа"))
        session.commit()
    book = workbook(env.client.get(prefix(first) + "/api/clients/export.xlsx"))
    sheet = book["Клиенты"]
    assert book.sheetnames == ["Клиенты", "О выгрузке"]
    assert sheet.max_row == 68 and sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:N68"
    rows = {row[0].value: row for row in list(sheet.rows)[1:]}
    assert rows["99"][1].data_type == "s" and rows["99"][1].value.startswith(
        "=HYPERLINK"
    )
    assert rows["99"][4].value == "+79990000000"
    assert rows["99"][4].data_type == "s" and rows["99"][0].data_type == "s"
    assert rows["99"][8].value == when and rows["99"][10].value == "Да"
    assert rows["100"][1].value == "ИванПётр" and rows["100"][2].data_type == "s"
    assert book["О выгрузке"]["B1"].value == first.name
    book.close()
    filtered = workbook(
        env.client.get(
            prefix(first) + "/api/clients/export.xlsx",
            params={"q": "Анна", "contacted": "true"},
        )
    )
    assert filtered["Клиенты"].max_row == 34
    filtered.close()
    other = workbook(env.client.get(prefix(second) + "/api/clients/export.xlsx"))
    assert (
        other["Клиенты"].max_row == 2
        and other["Клиенты"]["B2"].value == "Другая группа"
    )
    other.close()
    empty = workbook(
        env.client.get(prefix(second) + "/api/clients/export.xlsx?q=НетТакого")
    )
    assert empty["Клиенты"].max_row == 1
    empty.close()
    assert not env.network.calls


def test_excel_requires_owner_and_unknown_project_denied(project_env):
    env = project_env
    assert env.client.get("/admin/api/clients/export.xlsx").status_code == 401
    login(env.client)
    assert env.client.get("/p/999/admin/api/clients/export.xlsx").status_code == 404


def test_video_credential_encrypted_preserved_clear_and_no_leak(two_projects):
    env, first, second = two_projects
    private = "sensitive-video-credential"
    first = configure(env.client, first, video_token=private)
    assert first.video_token == private
    assert configure(env.client, first, video_token="").video_token == private
    assert projects.get_project(second.id).video_token == ""
    with env.engine.connect() as connection:
        assert private not in str(
            connection.execute(text("SELECT * FROM projects")).all()
        )
    assert private not in env.client.get("/projects").text
    assert private not in json.dumps(projects.public_project(first))
    with projects.project_scope(first):
        assert private not in json.dumps(config.get_settings().model_dump())
    assert configure(env.client, first, clear_video_token="1").video_token == ""
    assert not env.network.calls


def test_video_public_reply_uses_own_token_as_group_and_no_duplicates(two_projects):
    env, first, second = two_projects
    for original in (first, second):
        project = configure(env.client, original, video_token=f"video-{original.id}")
        campaign(project, delivery_mode="chat_invite")
        callback(env.client, project, video(project))
        callback(env.client, project, video(project))
        callback(env.client, project, video(project, number=2))
        with session_for(project) as session:
            assert session.query(db.PendingGift).count() == 1
        assert status(project) == "invite_duplicate"
    replies = [c for c in env.network.calls if c["method"] == "video.createComment"]
    assert len(replies) == 2
    for call, project in zip(replies, (first, second)):
        params = call["params"]
        assert params["access_token"] == f"video-{project.id}"
        assert params["from_group"] == "1" and params["owner_id"] == str(
            -project.group_id
        )
        assert params["video_id"] == "55" and params["reply_to_comment"] == "1"
        assert params["guid"]
    callback(env.client, first, message(first, text="Подарок"))
    sends = [c for c in env.network.calls if c["method"] == "messages.send"]
    assert len(sends) == 1 and sends[0]["params"]["access_token"] == first.token
    assert "GIFT" in sends[0]["params"]["message"]


@pytest.mark.parametrize("allowed", [True, False])
def test_video_without_user_token_fallback_and_eventual_claim(two_projects, allowed):
    env, first, _ = two_projects
    campaign(first, delivery_mode="chat_invite")
    env.network.replies["messages.isMessagesFromGroupAllowed"] = {
        "response": {"is_allowed": int(allowed)}
    }
    callback(env.client, first, video(first))
    assert status(first) == ("video_invited_dm" if allowed else "video_waiting_chat")
    calls = list(env.network.calls)
    callback(env.client, first, video(first))
    callback(env.client, first, video(first, number=2))
    assert env.network.calls == calls
    assert not any(c["method"] == "video.createComment" for c in calls)
    sends = [c for c in calls if c["method"] == "messages.send"]
    assert len(sends) == int(allowed)
    if allowed:
        assert "Получить подарок" in sends[0]["params"]["keyboard"]
    callback(env.client, first, message(first, text="Подарок"))
    # Later duplicate comments have separate statuses; the original gets sent.
    with session_for(first) as session:
        assert session.query(db.PendingGift).one().status == "sent"
        assert (
            session.query(db.ProcessedComment).filter_by(comment_id=1).one().status
            == "sent"
        )


@pytest.mark.parametrize("reason", ["unsubscribed", "permission_race"])
def test_video_respects_optout_and_901_without_losing_gift(two_projects, reason):
    env, first, _ = two_projects
    campaign(first, delivery_mode="chat_invite")
    if reason == "unsubscribed":
        with session_for(first) as session:
            session.add(db.Client(user_id=77, unsubscribed=True))
            session.commit()
    else:
        env.network.replies["messages.send"] = {
            "error": {"error_code": 901, "error_msg": "No permission"}
        }
    callback(env.client, first, video(first))
    assert status(first) == "video_waiting_chat"
    with session_for(first) as session:
        assert session.query(db.PendingGift).one().status == "pending"
    if reason == "unsubscribed":
        assert env.network.calls == []


def test_video_api_error_retains_gift_no_personal_fallback_and_masks_token(
    two_projects,
):
    env, first, _ = two_projects
    private = "secret-personal-key"
    first = configure(env.client, first, video_token=private)
    campaign(first, delivery_mode="chat_invite")
    env.network.replies["video.createComment"] = {
        "error": {"error_code": 5, "error_msg": "Invalid " + private}
    }
    callback(env.client, first, video(first))
    callback(env.client, first, video(first))
    assert len(env.network.calls) == 1
    with session_for(first) as session:
        row = session.query(db.ProcessedComment).one()
        assert row.status == "failed" and private not in row.error
        assert session.query(db.PendingGift).count() == 1


def test_video_boundary_and_paused_project_no_calls(two_projects):
    env, first, second = two_projects
    first = configure(env.client, first, video_token="video-key")
    campaign(first, delivery_mode="chat_invite")
    normalized = normalize_comment(video(second), second.group_id)
    with projects.project_scope(first), pytest.raises(ValueError):
        asyncio.run(vk_api.reply_to_video_comment(normalized, "text", "guid"))
    paused = configure(env.client, first, enabled=False)
    callback(env.client, paused, video(paused))
    with projects.project_scope(first), pytest.raises(vk_api.VkApiError):
        asyncio.run(
            vk_api.reply_to_video_comment(
                normalize_comment(video(first), first.group_id), "text", "guid"
            )
        )
    assert env.network.calls == []


def test_plus_additive_migration_preserves_campaigns(two_projects):
    _, first, second = two_projects
    for project in (first, second):
        campaign(project, plus_words="")
        engine = db.get_project_engine(project)
        with engine.begin() as connection:
            connection.execute(text("ALTER TABLE campaigns DROP COLUMN plus_words"))
        db.init_db(engine, migration_owner_group_id=project.group_id)
        db.init_db(engine, migration_owner_group_id=project.group_id)
        with session_for(project) as session:
            row = session.query(db.Campaign).one()
            assert row.plus_words == "" and row.promo_code == "GIFT"
